"""Excel workbooks.

Custody and clearing statements arrive as workbooks built for humans, not for
parsers. They open with a cover sheet carrying a disclaimer, the data sheet
leads with a title banner and a blank spacer, headers are sometimes merged
across columns, and the last row is a total that is not a holding.

A parser that takes sheet zero, row zero produces a plausible table of
disclaimer text. A parser that takes the last row as data inflates the
reconciliation by one and produces a break that looks real. Both failures are
quiet, which is why each is handled explicitly here.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from inspector.parsers.base import ParsedTable, ParseError

# Sheets whose names suggest they are not the data.
_NON_DATA_SHEET_TOKENS = (
    "cover",
    "disclaimer",
    "notes",
    "legend",
    "glossary",
    "index",
    "contents",
    "summary",
    "about",
    "readme",
)

_TOTAL_MARKERS = frozenset(
    {"total", "totals", "grandtotal", "subtotal", "sum", "count", "endofreport", "eof", "end"}
)

_XLSX_MAGIC = b"PK\x03\x04"
_MAX_HEADER_SEARCH_ROWS = 25


def _norm(value: Any) -> str:
    return "".join(ch for ch in str(value).lower() if ch.isalnum())


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _looks_numeric(value: Any) -> bool:
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip().replace(",", "").replace("$", "")
    if not text:
        return False
    try:
        float(text)
    except ValueError:
        return False
    return True


class ExcelParser:
    """Parser for .xlsx workbooks."""

    format_name = "EXCEL"

    def sniff(self, head: bytes, name: str) -> float:
        # xlsx is a zip archive. The magic bytes are decisive and the
        # extension is not, which is the whole point of FR-17: a workbook
        # renamed .csv is still a workbook.
        if not head.startswith(_XLSX_MAGIC):
            return 0.0
        score = 0.9
        if name.lower().endswith((".xlsx", ".xlsm")):
            score = 0.97
        return score

    def parse(self, path: Path) -> ParsedTable:
        import zipfile

        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            # read_only would be faster, but merged-cell ranges are not
            # available in that mode and forward-filling merged headers is
            # exactly what this parser exists to do.
            wb = load_workbook(path, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
            # A file with an .xlsx name that is not a zip archive is a
            # malformed arrival, not a crash. It must reject cleanly so the
            # error taxonomy can route it (MALFORMED, no retry).
            raise ParseError(f"not a readable workbook: {exc}") from exc

        try:
            sheet_name = self._choose_sheet(wb)
            sheet = wb[sheet_name]
            grid = self._grid(sheet)
        finally:
            wb.close()

        if not grid:
            raise ParseError(f"sheet {sheet_name!r} contains no data")

        header_idx = self._find_header_row(grid)
        headers = self._dedupe(grid[header_idx])
        if not headers:
            raise ParseError("no header row could be identified")

        warnings: list[str] = []
        if len(wb.sheetnames) > 1:
            warnings.append(f"workbook has {len(wb.sheetnames)} sheets; selected {sheet_name!r}")
        if header_idx > 0:
            warnings.append(f"skipped {header_idx} row(s) above the header")

        rows: list[dict[str, str]] = []
        dropped = 0
        for raw in grid[header_idx + 1 :]:
            if all(_is_blank(cell) for cell in raw):
                continue
            if self._is_total_row(raw):
                dropped += 1
                continue
            rows.append(
                {
                    h: ("" if _is_blank(raw[i]) else str(raw[i]).strip()) if i < len(raw) else ""
                    for i, h in enumerate(headers)
                }
            )
        if dropped:
            warnings.append(f"excluded {dropped} summary row(s)")

        return ParsedTable(
            headers=headers,
            rows=rows,
            evidence={
                "sheet": sheet_name,
                "sheet_count": str(len(wb.sheetnames)),
                "header_row_index": str(header_idx),
            },
            warnings=warnings,
        )

    def _choose_sheet(self, wb: Any) -> str:
        """Pick the data sheet, not the cover sheet.

        Scored rather than assumed: a sheet named "Cover" is demoted, and a
        sheet with more populated rows and columns is preferred. Falling back
        to index zero is how the disclaimer becomes the dataset.
        """
        best_name: str = str(wb.sheetnames[0])
        best_score = float("-inf")
        for name in wb.sheetnames:
            sheet = wb[name]
            rows = sheet.max_row or 0
            cols = sheet.max_column or 0
            score = min(rows, 5000) * 1.0 + min(cols, 100) * 5.0
            if any(token in _norm(name) for token in _NON_DATA_SHEET_TOKENS):
                score -= 5000.0
            if rows <= 1 or cols <= 1:
                score -= 1000.0
            if score > best_score:
                best_name, best_score = str(name), score
        return best_name

    def _grid(self, sheet: Any) -> list[list[Any]]:
        """Read cells, forward-filling any merged ranges.

        A merged header spanning three columns reports its value only in the
        top-left cell; the others come back as None. Left unfilled, two of the
        three columns lose their names and the data under them is orphaned.
        """
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        merged_cells = getattr(sheet, "merged_cells", None)
        for merged in list(getattr(merged_cells, "ranges", [])):
            min_row, min_col = merged.min_row - 1, merged.min_col - 1
            if min_row >= len(grid) or min_col >= len(grid[min_row]):
                continue
            value = grid[min_row][min_col]
            for r in range(merged.min_row - 1, merged.max_row):
                for c in range(merged.min_col - 1, merged.max_col):
                    if r < len(grid) and c < len(grid[r]):
                        grid[r][c] = value
        # Trim wholly empty trailing rows, which openpyxl reports generously.
        while grid and all(_is_blank(cell) for cell in grid[-1]):
            grid.pop()
        return grid

    def _find_header_row(self, grid: list[list[Any]]) -> int:
        """Same scoring as the delimited parser: labels, not data."""
        widths = [sum(1 for cell in row if not _is_blank(cell)) for row in grid]
        if not widths:
            return 0
        modal_width = max(set(widths), key=lambda w: (widths.count(w), w))

        best_idx, best_score = 0, float("-inf")
        for idx in range(min(len(grid), _MAX_HEADER_SEARCH_ROWS)):
            filled = [cell for cell in grid[idx] if not _is_blank(cell)]
            if not filled:
                continue
            non_numeric = sum(1 for cell in filled if not _looks_numeric(cell))
            score = 2.0 * (non_numeric / len(filled))
            score += 1.5 if len(filled) == modal_width else 0.0
            score += 1.0 * (len(filled) / max(len(grid[idx]), 1))
            score -= 0.15 * idx
            if score > best_score:
                best_idx, best_score = idx, score
        return best_idx

    def _is_total_row(self, row: list[Any]) -> bool:
        first = next((cell for cell in row if not _is_blank(cell)), "")
        return _norm(first) in _TOTAL_MARKERS

    def _dedupe(self, raw_headers: list[Any]) -> list[str]:
        seen: dict[str, int] = {}
        out: list[str] = []
        trailing_blanks = 0
        for i, cell in enumerate(raw_headers):
            if _is_blank(cell):
                trailing_blanks += 1
                name = f"column_{i + 1}"
            else:
                trailing_blanks = 0
                name = str(cell).strip()
            if name in seen:
                seen[name] += 1
                name = f"{name}__{seen[name]}"
            else:
                seen[name] = 0
            out.append(name)
        # Drop placeholder columns that trail off the end of the header row.
        while out and trailing_blanks > 0 and out[-1].startswith("column_"):
            out.pop()
            trailing_blanks -= 1
        return out
